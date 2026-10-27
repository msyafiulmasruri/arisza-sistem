"""Backend sistem web Arisza (FastAPI di Vercel).

Tugasnya tipis: menerima impor, mencatat permintaan di PostgreSQL, memanggil layanan AI di
Hugging Face Space, lalu membaca hasil dari basis data. Model tidak dijalankan di sini.
"""
import csv
import hmac
import io
import os
from contextlib import contextmanager

import httpx
import psycopg
from fastapi import Depends, FastAPI, File, Form, Header, HTTPException, Query, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from psycopg.rows import dict_row

DATABASE_URL = os.getenv("DATABASE_URL", "")
ADMIN_PASSWORD = os.getenv("ADMIN_PASSWORD", "")
SPACE_URL = os.getenv("SPACE_URL", "").rstrip("/")
SPACE_API_KEY = os.getenv("SPACE_API_KEY", "")
HF_SPACE_TOKEN = os.getenv("HF_SPACE_TOKEN", "")
FRONTEND_ORIGINS = [o.strip() for o in os.getenv("FRONTEND_ORIGIN", "http://localhost:5173").split(",") if o.strip()]
BATAS_MACET_MENIT = int(os.getenv("BATAS_MACET_MENIT", "30"))
MAKS_BARIS_IMPOR = int(os.getenv("MAKS_BARIS_IMPOR", "5000"))

KOLOM_TEKS = {"text", "teks", "komentar", "ulasan", "review", "comment", "tanggapan", "isi"}
KOLOM_MEDIA = {"media", "sumber", "platform", "source"}
TAHAP = ["positive", "neutral", "negative", "saran"]
LABEL = {"positive", "neutral", "negative"}

app = FastAPI(title="Arisza Backend")
app.add_middleware(
    CORSMiddleware, allow_origins=FRONTEND_ORIGINS, allow_methods=["*"], allow_headers=["*"],
)


# ---------- dasar ----------

def admin(x_admin_key: str = Header(default="")):
    if not ADMIN_PASSWORD:
        raise HTTPException(503, "ADMIN_PASSWORD belum diatur di backend.")
    if not hmac.compare_digest(x_admin_key.encode(), ADMIN_PASSWORD.encode()):
        raise HTTPException(401, "Kata sandi salah.")


@contextmanager
def konek():
    if not DATABASE_URL:
        raise HTTPException(503, "DATABASE_URL belum diatur di backend.")
    conn = psycopg.connect(
        DATABASE_URL, row_factory=dict_row, autocommit=True, prepare_threshold=None, connect_timeout=10
    )
    try:
        yield conn
    finally:
        conn.close()


def panggil_space(path, payload):
    if not SPACE_URL or not SPACE_API_KEY:
        raise RuntimeError("SPACE_URL atau SPACE_API_KEY belum diatur di backend.")
    headers = {"X-API-Key": SPACE_API_KEY}
    if HF_SPACE_TOKEN:
        headers["Authorization"] = f"Bearer {HF_SPACE_TOKEN}"
    # Space yang tertidur butuh waktu bangun, jadi batas waktunya longgar.
    r = httpx.post(f"{SPACE_URL}{path}", json=payload, headers=headers, timeout=90)
    r.raise_for_status()


def tandai_macet(conn):
    """Pekerjaan yang tidak bergerak lama (Space restart atau tidur) ditandai gagal agar bisa diulang."""
    pesan = "Proses terhenti karena melewati batas waktu. Gunakan tombol ulangi."
    ids = [
        r["id"]
        for r in conn.execute(
            "update interpretation_runs set status='failed', error=%s, updated_at=now() "
            "where status in ('queued','running') and updated_at < now() - make_interval(mins => %s) returning id",
            (pesan, BATAS_MACET_MENIT),
        ).fetchall()
    ]
    if ids:
        conn.execute(
            "update interpretation_stages set status='failed', error=%s where run_id = any(%s) and status='running'",
            (pesan, ids),
        )
    conn.execute(
        "update import_batches set status='failed', error=%s, updated_at=now() "
        "where status in ('queued','running') and updated_at < now() - make_interval(mins => %s)",
        (pesan, BATAS_MACET_MENIT),
    )


def muat_run(conn, run_id):
    run = conn.execute("select * from interpretation_runs where id=%s", (run_id,)).fetchone()
    if run is None:
        return None
    tahap = conn.execute("select * from interpretation_stages where run_id=%s", (run_id,)).fetchall()
    run["tahap"] = sorted(tahap, key=lambda t: TAHAP.index(t["stage"]))
    return run


# ---------- umum ----------

@app.get("/api/health")
def health():
    return {"status": "ok"}


@app.get("/api/auth", dependencies=[Depends(admin)])
def auth():
    return {"ok": True}


# ---------- impor ----------

def baca_berkas(nama, isi):
    nama = (nama or "").lower()
    if nama.endswith(".csv"):
        teks = isi.decode("utf-8-sig", errors="replace")
        try:
            dialek = csv.Sniffer().sniff(teks[:4096], delimiters=",;\t")
        except csv.Error:
            dialek = csv.excel
        return [list(r) for r in csv.reader(io.StringIO(teks), dialek)]
    if nama.endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(isi), read_only=True, data_only=True)
        return [list(r) for r in wb.active.iter_rows(values_only=True)]
    raise HTTPException(400, "Format berkas harus .csv atau .xlsx.")


def rapikan_media(nama):
    return " ".join(str(nama).split())


@app.post("/api/import", dependencies=[Depends(admin)], status_code=202)
async def impor(file: UploadFile = File(...), media_bawaan: str = Form("")):
    baris = baca_berkas(file.filename, await file.read())
    if len(baris) < 2:
        raise HTTPException(400, "Berkas kosong atau hanya berisi judul kolom.")
    judul = [str(h or "").strip().lower() for h in baris[0]]
    kt = next((i for i, h in enumerate(judul) if h in KOLOM_TEKS), None)
    km = next((i for i, h in enumerate(judul) if h in KOLOM_MEDIA), None)
    if kt is None:
        raise HTTPException(
            400, "Kolom teks tidak ditemukan. Beri judul kolom salah satu dari: " + ", ".join(sorted(KOLOM_TEKS)) + "."
        )
    media_bawaan = rapikan_media(media_bawaan)
    if km is None and not media_bawaan:
        raise HTTPException(400, "Berkas tidak punya kolom media. Isi kolom 'media' atau tulis nama media pada formulir.")

    with konek() as conn:
        ada = {r["media"].lower(): r["media"] for r in conn.execute("select distinct media from responses").fetchall()}
        data = []
        for r in baris[1:]:
            teks = str(r[kt]).strip() if kt < len(r) and r[kt] is not None else ""
            media = rapikan_media(r[km]) if km is not None and km < len(r) and r[km] else media_bawaan
            if not teks or not media:
                continue
            media = ada.setdefault(media.lower(), media)
            data.append((media, teks))
        if not data:
            raise HTTPException(400, "Tidak ada baris yang berisi teks dan media.")
        if len(data) > MAKS_BARIS_IMPOR:
            raise HTTPException(400, f"Maksimal {MAKS_BARIS_IMPOR} baris per berkas; bagi berkas menjadi beberapa bagian.")

        batch = conn.execute(
            "insert into import_batches (filename, total) values (%s, %s) returning id", (file.filename, len(data))
        ).fetchone()
        with conn.cursor() as cur:
            cur.executemany(
                "insert into responses (batch_id, media, text) values (%s, %s, %s)",
                [(batch["id"], m, t) for m, t in data],
            )
        kirim_batch(conn, batch["id"])
        return conn.execute("select * from import_batches where id=%s", (batch["id"],)).fetchone()


def kirim_batch(conn, batch_id):
    try:
        panggil_space("/jobs/classify", {"batch_id": batch_id})
    except Exception as exc:
        conn.execute(
            "update import_batches set status='failed', error=%s, updated_at=now() where id=%s",
            (f"Gagal menghubungi layanan AI: {exc}"[:500], batch_id),
        )


@app.get("/api/batches", dependencies=[Depends(admin)])
def daftar_batch():
    with konek() as conn:
        tandai_macet(conn)
        return conn.execute(
            "select b.*, count(r.id) filter (where r.status<>'pending') as diproses "
            "from import_batches b left join responses r on r.batch_id=b.id "
            "group by b.id order by b.id desc limit 10"
        ).fetchall()


@app.post("/api/batches/{batch_id}/retry", dependencies=[Depends(admin)], status_code=202)
def ulangi_batch(batch_id: int):
    with konek() as conn:
        b = conn.execute("select status from import_batches where id=%s", (batch_id,)).fetchone()
        if b is None:
            raise HTTPException(404, "Batch tidak ditemukan.")
        if b["status"] != "failed":
            raise HTTPException(409, "Hanya batch yang gagal yang bisa diulang.")
        conn.execute("update import_batches set status='queued', error=null, updated_at=now() where id=%s", (batch_id,))
        kirim_batch(conn, batch_id)
        return conn.execute("select * from import_batches where id=%s", (batch_id,)).fetchone()


# ---------- dasbor ----------

@app.get("/api/summary", dependencies=[Depends(admin)])
def ringkasan():
    with konek() as conn:
        tandai_macet(conn)
        media = conn.execute(
            "select media, "
            "count(*) filter (where status='ok') as total, "
            "count(*) filter (where label='positive') as positive, "
            "count(*) filter (where label='neutral') as neutral, "
            "count(*) filter (where label='negative') as negative, "
            "count(*) filter (where status='pending') as pending, "
            "count(*) filter (where status='ok' and (clean_text is null or wc_text is null)) as perlu_lengkap "
            "from responses group by media order by media"
        ).fetchall()
        terbaru = {
            r["media"]: r
            for r in conn.execute(
                "select distinct on (media) id, media, status, error from interpretation_runs order by media, id desc"
            ).fetchall()
        }
        selesai = {
            r["media"]: r
            for r in conn.execute(
                "select distinct on (r.media) r.id, r.media, r.n_texts, r.finished_at, "
                "(select sum(s.duration) from interpretation_stages s where s.run_id = r.id) as total_detik "
                "from interpretation_runs r where r.status='done' order by r.media, r.id desc"
            ).fetchall()
        }
        for m in media:
            t, s = terbaru.get(m["media"]), selesai.get(m["media"])
            m["interpretasi"] = {
                "status": t["status"] if t else None,
                "run_id": t["id"] if t else None,
                "selesai_pada": s["finished_at"] if s else None,
                "total_detik": float(s["total_detik"]) if s and s["total_detik"] is not None else None,
                "n_teks": s["n_texts"] if s else None,
                "ada_hasil": s is not None,
                # data bertambah sejak interpretasi terakhir
                "basi": bool(s) and s["n_texts"] != m["total"],
            }
        return media


@app.get("/api/media/{media}/wordclouds", dependencies=[Depends(admin)])
def info_wordcloud(media: str):
    """Keterangan Word Cloud yang tersedia (gambarnya diambil lewat endpoint terpisah)."""
    with konek() as conn:
        return conn.execute(
            "select label, n_texts, updated_at from wordclouds where media=%s order by label", (media,)
        ).fetchall()


@app.get("/api/media/{media}/wordcloud/{label}", dependencies=[Depends(admin)])
def gambar_wordcloud(media: str, label: str):
    if label not in LABEL:
        raise HTTPException(404, "Sentimen tidak dikenal.")
    with konek() as conn:
        r = conn.execute("select image from wordclouds where media=%s and label=%s", (media, label)).fetchone()
    if r is None:
        raise HTTPException(404, "Word Cloud belum tersedia.")
    return Response(bytes(r["image"]), media_type="image/png", headers={"Cache-Control": "no-store"})


@app.post("/api/maintenance/lengkapi", dependencies=[Depends(admin)], status_code=202)
def lengkapi_data_lama():
    """Minta Space mengisi teks preprocessing dan Word Cloud untuk data lama (label tidak diubah)."""
    try:
        panggil_space("/jobs/complete", {})
    except Exception as exc:
        raise HTTPException(502, f"Gagal menghubungi layanan AI: {exc}"[:300])
    return {"accepted": True}


# ---------- tabel data ----------

def _saring(media, label, q):
    kondisi, params = [], {}
    if media:
        kondisi.append("media = %(media)s")
        params["media"] = media
    if label:
        if label == "empty":
            kondisi.append("status = 'empty'")
        elif label in LABEL:
            kondisi.append("label = %(label)s")
            params["label"] = label
        else:
            raise HTTPException(400, "Sentimen harus positive, neutral, negative, atau empty.")
    if q:
        kondisi.append("(text ilike %(q)s or clean_text ilike %(q)s)")
        params["q"] = "%" + q.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
    return ("where " + " and ".join(kondisi)) if kondisi else "", params


@app.get("/api/responses", dependencies=[Depends(admin)])
def tabel(
    media: str = "", label: str = "", q: str = Query("", max_length=100),
    halaman: int = Query(1, ge=1), per_halaman: int = Query(25, ge=5, le=100),
):
    where, params = _saring(media.strip(), label.strip(), q.strip())
    with konek() as conn:
        total = conn.execute(f"select count(*) as n from responses {where}", params).fetchone()["n"]
        baris = conn.execute(
            "select id, media, text, clean_text, label, confidence, status from responses "
            f"{where} order by id limit %(lim)s offset %(off)s",
            {**params, "lim": per_halaman, "off": (halaman - 1) * per_halaman},
        ).fetchall()
    return {"total": total, "halaman": halaman, "per_halaman": per_halaman, "baris": baris}


@app.get("/api/responses/export", dependencies=[Depends(admin)])
def ekspor(media: str = "", label: str = "", q: str = Query("", max_length=100)):
    where, params = _saring(media.strip(), label.strip(), q.strip())
    with konek() as conn:
        baris = conn.execute(
            "select media, text, clean_text, label, confidence from responses "
            f"{where} order by id limit 20000",
            params,
        ).fetchall()
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["media", "teks_awal", "teks_preprocessing", "sentimen", "confidence"])
    for b in baris:
        w.writerow([b["media"], b["text"], b["clean_text"] or "", b["label"] or "", "" if b["confidence"] is None else round(b["confidence"], 4)])
    return Response(
        "\ufeff" + out.getvalue(), media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="tanggapan.csv"'},
    )


# ---------- interpretasi ----------

@app.get("/api/media/{media}/interpretation", dependencies=[Depends(admin)])
def interpretasi(media: str):
    with konek() as conn:
        tandai_macet(conn)
        t = conn.execute("select id from interpretation_runs where media=%s order by id desc limit 1", (media,)).fetchone()
        s = conn.execute(
            "select id from interpretation_runs where media=%s and status='done' order by id desc limit 1", (media,)
        ).fetchone()
        jumlah = conn.execute(
            "select count(*) filter (where status='ok') as total, count(*) filter (where status='pending') as pending "
            "from responses where media=%s",
            (media,),
        ).fetchone()
        return {
            "terbaru": muat_run(conn, t["id"]) if t else None,
            "terakhir_selesai": muat_run(conn, s["id"]) if s else None,
            "jumlah_teks": jumlah["total"],
            "menunggu_prediksi": jumlah["pending"],
        }


def kirim_run(conn, run_id):
    try:
        panggil_space("/jobs/interpret", {"run_id": run_id})
    except Exception as exc:
        conn.execute(
            "update interpretation_runs set status='failed', error=%s, updated_at=now() where id=%s",
            (f"Gagal menghubungi layanan AI: {exc}"[:500], run_id),
        )


@app.post("/api/media/{media}/interpret", dependencies=[Depends(admin)], status_code=202)
def jalankan_interpretasi(media: str):
    with konek() as conn:
        tandai_macet(conn)
        jumlah = conn.execute(
            "select count(*) filter (where status='ok') as total, count(*) filter (where status='pending') as pending "
            "from responses where media=%s",
            (media,),
        ).fetchone()
        if jumlah["total"] == 0:
            raise HTTPException(400, "Belum ada tanggapan berlabel pada media ini.")
        if jumlah["pending"] > 0:
            raise HTTPException(409, "Prediksi sentimen untuk data impor masih berjalan. Tunggu sampai selesai.")
        aktif = conn.execute(
            "select id from interpretation_runs where media=%s and status in ('queued','running')", (media,)
        ).fetchone()
        if aktif:
            raise HTTPException(409, "Interpretasi untuk media ini sedang diproses.")
        run = conn.execute("insert into interpretation_runs (media) values (%s) returning id", (media,)).fetchone()
        with conn.cursor() as cur:
            cur.executemany(
                "insert into interpretation_stages (run_id, stage) values (%s, %s)", [(run["id"], t) for t in TAHAP]
            )
        kirim_run(conn, run["id"])
        return muat_run(conn, run["id"])


@app.post("/api/runs/{run_id}/retry", dependencies=[Depends(admin)], status_code=202)
def ulangi_run(run_id: int):
    with konek() as conn:
        run = conn.execute("select media, status from interpretation_runs where id=%s", (run_id,)).fetchone()
        if run is None:
            raise HTTPException(404, "Pekerjaan tidak ditemukan.")
        if run["status"] != "failed":
            raise HTTPException(409, "Hanya pekerjaan yang gagal yang bisa diulang.")
        aktif = conn.execute(
            "select id from interpretation_runs where media=%s and status in ('queued','running')", (run["media"],)
        ).fetchone()
        if aktif:
            raise HTTPException(409, "Interpretasi untuk media ini sedang diproses.")
        conn.execute(
            "update interpretation_stages set status='waiting', error=null where run_id=%s and status in ('failed','running')",
            (run_id,),
        )
        conn.execute("update interpretation_runs set status='queued', error=null, updated_at=now() where id=%s", (run_id,))
        kirim_run(conn, run_id)
        return muat_run(conn, run_id)
